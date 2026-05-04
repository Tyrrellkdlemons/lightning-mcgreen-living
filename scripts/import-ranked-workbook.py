#!/usr/bin/env python3
"""
Import ranked apartment workbook rows into JSON used by the apply engine.

Usage:
  python scripts/import-ranked-workbook.py
  python scripts/import-ranked-workbook.py --input "C:\\path\\file.xlsx"
  python scripts/import-ranked-workbook.py --input "C:\\path\\file.xlsx" --output "src/lib/data/generated/workbook-ranked-listings.json"
"""

from __future__ import annotations

import argparse
import json
import re
import time
import urllib.parse
import urllib.request
from pathlib import Path

from openpyxl import load_workbook


def as_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace(",", "")
    if s == "" or s.lower().startswith("verify"):
        return None
    try:
        return float(s)
    except Exception:
        return None


def as_int(v):
    n = as_num(v)
    return int(n) if n is not None else None


def as_str(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def slug(s: str):
    s = (s or "").lower()
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s[:80]


def is_url(v):
    if not isinstance(v, str):
        return False
    return bool(re.match(r"^https?://", v.strip(), re.I))


def ducky_url(query: str):
    return "https://duckduckgo.com/?q=" + urllib.parse.quote("!ducky " + query)


def infer_platform(url_blob: str | None, mgmt: str | None, flag: str | None):
    b = f"{url_blob or ''} {mgmt or ''} {flag or ''}".lower()
    if "appfolio" in b:
        return "AppFolio"
    if "on-site" in b or "on-site.com" in b:
        return "On-Site"
    if "rentcafe" in b:
        return "RentCafe"
    if "entrata" in b:
        return "Entrata"
    if "realpage" in b:
        return "RealPage"
    if "knock" in b or "g5" in b:
        return "G5/Knock"
    return "Other"


def unique_links(links):
    out = []
    seen = set()
    for link in links:
        url = link.get("url")
        if not is_url(url):
            continue
        if url in seen:
            continue
        seen.add(url)
        out.append(link)
    return out


def pick_primary_url(*values):
    for v in values:
        if is_url(v):
            return v
    return None


def host_for_url(url: str | None):
    if not is_url(url):
        return None
    return urllib.parse.urlparse(url).netloc.lower().replace("www.", "")


PLATFORM_HOSTS = {
    "AppFolio": "appfolio.com",
    "On-Site": "on-site.com",
    "RentCafe": "rentcafe.com",
    "Entrata": "entrata.com",
    "RealPage": "realpage.com",
    "G5/Knock": "knockcrm.com",
}


def map_links(address: str | None, city: str | None, county: str | None, lat: float | None, lng: float | None):
    base_q = ", ".join([x for x in [address, city, county, "California"] if x])
    if lat is not None and lng is not None:
        return {
            "lat": lat,
            "lng": lng,
            "openstreetmap_url": f"https://www.openstreetmap.org/?mlat={lat:.6f}&mlon={lng:.6f}#map=16/{lat:.6f}/{lng:.6f}",
            "openstreetmap_embed_url": f"https://www.openstreetmap.org/export/embed.html?bbox={lng-0.01:.6f}%2C{lat-0.01:.6f}%2C{lng+0.01:.6f}%2C{lat+0.01:.6f}&layer=mapnik&marker={lat:.6f}%2C{lng:.6f}",
            "google_maps_reference_url": f"https://www.google.com/maps?q={lat:.6f},{lng:.6f}",
        }
    q = urllib.parse.quote(base_q or "Southern California apartment")
    return {
        "lat": None,
        "lng": None,
        "openstreetmap_url": f"https://www.openstreetmap.org/search?query={q}",
        "openstreetmap_embed_url": None,
        "google_maps_reference_url": f"https://www.google.com/maps/search/?api=1&query={q}",
    }


def geocode_nominatim(query: str, cache: dict):
    if not query:
        return None, None
    key = query.strip().lower()
    if key in cache:
        pair = cache[key] or {}
        return pair.get("lat"), pair.get("lng")
    url = "https://nominatim.openstreetmap.org/search?" + urllib.parse.urlencode(
        {"q": query, "format": "jsonv2", "limit": 1}
    )
    req = urllib.request.Request(
        url,
        headers={"User-Agent": "lightning-mcgreen-living/1.0 (workbook geocode enrichment)"},
    )
    lat = None
    lng = None
    try:
        with urllib.request.urlopen(req, timeout=20) as r:
            body = r.read().decode("utf-8")
            data = json.loads(body)
            if data:
                lat = float(data[0]["lat"])
                lng = float(data[0]["lon"])
    except Exception:
        lat = None
        lng = None
    cache[key] = {"lat": lat, "lng": lng}
    # Keep a polite request cadence for public Nominatim.
    time.sleep(1.05)
    return lat, lng


def parse_args():
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--input",
        default=r"C:\Users\TKDL\Downloads\socal_apartment_ranked_research_workbook_ENHANCED_LINKED.xlsx",
        help="Workbook path",
    )
    parser.add_argument(
        "--output",
        default=r"src/lib/data/generated/workbook-ranked-listings.json",
        help="Output JSON path",
    )
    parser.add_argument(
        "--geocode-cache",
        default=r"src/lib/data/generated/workbook-geocode-cache.json",
        help="Geocode cache JSON path",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    src = Path(args.input)
    out = Path(args.output)
    cache_path = Path(args.geocode_cache)
    out.parent.mkdir(parents=True, exist_ok=True)
    cache_path.parent.mkdir(parents=True, exist_ok=True)

    if not src.exists():
        raise SystemExit(f"Input workbook not found: {src}")

    geocode_cache = {}
    if cache_path.exists():
        try:
            geocode_cache = json.loads(cache_path.read_text(encoding="utf-8"))
        except Exception:
            geocode_cache = {}

    wb = load_workbook(src, data_only=True)
    if "Ranked Listings" not in wb.sheetnames:
        raise SystemExit('Sheet "Ranked Listings" not found in workbook.')
    ws = wb["Ranked Listings"]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    idx = {h: i for i, h in enumerate(headers)}

    rows = []
    for r in range(2, ws.max_row + 1):
        prop = as_str(ws.cell(r, idx["Property Name"] + 1).value)
        if not prop:
            continue
        city = as_str(ws.cell(r, idx["City"] + 1).value)
        county = as_str(ws.cell(r, idx["County"] + 1).value)
        address = as_str(ws.cell(r, idx["Address"] + 1).value)
        rank = as_int(ws.cell(r, idx["Rank"] + 1).value)
        match_percent = as_num(ws.cell(r, idx["Match %"] + 1).value)
        rent_low = as_num(ws.cell(r, idx["Rent Low"] + 1).value)
        rent_high = as_num(ws.cell(r, idx["Rent High"] + 1).value)
        screening_difficulty = as_num(ws.cell(r, idx["Screening Difficulty"] + 1).value)

        direct_unit_link = as_str(ws.cell(r, idx["Direct Unit Link"] + 1).value)
        floor_plan_link = as_str(ws.cell(r, idx["Floor Plan Link"] + 1).value)
        leasing_site = as_str(ws.cell(r, idx["Leasing Site"] + 1).value)
        source_url = as_str(ws.cell(r, idx["Source URL"] + 1).value)
        best_listing_url = as_str(ws.cell(r, idx["Best Listing URL"] + 1).value)
        floor_plan_url = as_str(ws.cell(r, idx["Floor Plan URL"] + 1).value)
        leasing_url = as_str(ws.cell(r, idx["Leasing URL"] + 1).value)
        source_url_clickable = as_str(ws.cell(r, idx["Source URL Clickable"] + 1).value)
        mgmt = as_str(ws.cell(r, idx["Mgmt/Leasing Company"] + 1).value)
        leasing_contact = as_str(ws.cell(r, idx["Leasing Contact"] + 1).value)
        greystar_flag = as_str(ws.cell(r, idx["Greystar/Snappt Flag"] + 1).value)

        url_blob = " ".join(
            [
                x
                for x in [
                    direct_unit_link,
                    floor_plan_link,
                    leasing_site,
                    source_url,
                    best_listing_url,
                    floor_plan_url,
                    leasing_url,
                    source_url_clickable,
                ]
                if x
            ]
        )
        platform = infer_platform(url_blob, mgmt, greystar_flag)

        primary_url = pick_primary_url(best_listing_url, leasing_url, source_url_clickable, source_url, leasing_site)
        source_host = host_for_url(primary_url) or "apartments.com"
        city_for_query = city or county or "California"
        listing_query = f'site:{source_host} "{prop}" "{city_for_query}" apartment'
        official_query = f'"{prop}" "{city_for_query}" apartment leasing'
        platform_host = PLATFORM_HOSTS.get(platform)
        platform_query = f'site:{platform_host} "{prop}" "{city_for_query}" apply' if platform_host else None

        direct_listing_url = ducky_url(listing_query)
        official_property_search_url = ducky_url(official_query)
        direct_platform_apply_url = ducky_url(platform_query) if platform_query else None
        zillow_url = ducky_url(f'site:zillow.com "{prop}" "{city_for_query}"')
        rent_url = ducky_url(f'site:rent.com "{prop}" "{city_for_query}"')
        realtor_url = ducky_url(f'site:realtor.com "{prop}" "{city_for_query}"')

        geocode_query = ", ".join([x for x in [address, city_for_query, "California"] if x])
        lat, lng = geocode_nominatim(geocode_query, geocode_cache)

        screening_precheck_note = None
        if greystar_flag:
            flag_l = greystar_flag.lower()
            has_strict_tag = bool(re.search(r"\b(greystar|snappt)\b", flag_l))
            explicit_negative = bool(
                re.search(r"\b(not|no|none|without)\b[^.]{0,32}\b(greystar|snappt)\b", flag_l)
            )
            if has_strict_tag and not explicit_negative:
                screening_precheck_note = "Strict-screening warning: this listing is flagged for Greystar/Snappt or similar checks."

        source_links = unique_links(
            [
                {"label": "Primary listing source", "url": primary_url, "kind": "primary"},
                {"label": "Direct listing deep link", "url": direct_listing_url, "kind": "deep-listing"},
                {"label": "Official property search", "url": official_property_search_url, "kind": "official-search"},
                {"label": "Platform apply deep link", "url": direct_platform_apply_url, "kind": "deep-apply"},
                {"label": "Zillow listing lookup", "url": zillow_url, "kind": "marketplace"},
                {"label": "Rent.com listing lookup", "url": rent_url, "kind": "marketplace"},
                {"label": "Realtor.com listing lookup", "url": realtor_url, "kind": "marketplace"},
                {"label": "Legacy best listing URL", "url": best_listing_url, "kind": "legacy"},
                {"label": "Legacy leasing URL", "url": leasing_url, "kind": "legacy"},
                {"label": "Legacy source URL", "url": source_url_clickable or source_url, "kind": "legacy"},
            ]
        )

        rec = {
            "id": f"wb-{rank or r}-{slug(prop)}-{slug(city or county or 'ca')}",
            "rank": rank,
            "match_percent": match_percent,
            "fit_tier": as_str(ws.cell(r, idx["Fit Tier"] + 1).value),
            "property_name": prop,
            "city": city,
            "county": county,
            "priority_city_match": as_str(ws.cell(r, idx["Priority City Match"] + 1).value),
            "neighborhood": as_str(ws.cell(r, idx["Neighborhood"] + 1).value),
            "address": address,
            "beds": as_str(ws.cell(r, idx["Beds"] + 1).value),
            "baths": as_str(ws.cell(r, idx["Baths"] + 1).value),
            "sqft_range": as_str(ws.cell(r, idx["Sq Ft / Range"] + 1).value),
            "rent_low": rent_low,
            "rent_high": rent_high,
            "unit_or_floor": as_str(ws.cell(r, idx["Unit # / Floor #"] + 1).value),
            "floor_plan": as_str(ws.cell(r, idx["Floor Plan"] + 1).value),
            "availability": as_str(ws.cell(r, idx["Availability"] + 1).value),
            "management_company": mgmt,
            "leasing_contact": leasing_contact,
            "screening_vendor": as_str(ws.cell(r, idx["Screening Vendor"] + 1).value),
            "screening_difficulty": screening_difficulty,
            "screening_difficulty_label": as_str(ws.cell(r, idx["Screening Difficulty Label"] + 1).value),
            "greystar_snappt_flag": greystar_flag,
            "action_status": as_str(ws.cell(r, idx["Action Status"] + 1).value),
            "score_notes": as_str(ws.cell(r, idx["Score Notes"] + 1).value),
            "best_feature": as_str(ws.cell(r, idx["Best Feature"] + 1).value),
            "missing_preferences": as_str(ws.cell(r, idx["Missing Preferences"] + 1).value),
            "questions_to_ask_leasing": as_str(ws.cell(r, idx["Questions To Ask Leasing"] + 1).value),
            "source_date": as_str(ws.cell(r, idx["Source Date"] + 1).value),
            "application_platform_inferred": platform,
            "screening_precheck_note": screening_precheck_note,
            "map": map_links(address, city, county, lat, lng),
            "direct_links": {
                "direct_listing_url": direct_listing_url,
                "official_property_search_url": official_property_search_url,
                "direct_platform_apply_url": direct_platform_apply_url,
                "zillow_url": zillow_url,
                "rent_url": rent_url,
                "realtor_url": realtor_url,
                "source_host": source_host,
            },
            "source_links": source_links,
            "links": {
                "direct_unit_link": direct_unit_link,
                "floor_plan_link": floor_plan_link,
                "leasing_site": leasing_site,
                "source_url": source_url,
                "best_listing_url": best_listing_url,
                "floor_plan_url": floor_plan_url,
                "leasing_url": leasing_url,
                "source_url_clickable": source_url_clickable,
            },
        }
        rows.append(rec)

    rows.sort(key=lambda x: (x["rank"] is None, x["rank"] or 9999, x["property_name"]))
    out.write_text(json.dumps(rows, indent=2), encoding="utf-8")
    cache_path.write_text(json.dumps(geocode_cache, indent=2), encoding="utf-8")
    print(f"Wrote {len(rows)} rows to {out}")


if __name__ == "__main__":
    main()
