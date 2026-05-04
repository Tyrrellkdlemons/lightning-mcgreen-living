#!/usr/bin/env python3
"""
import-workbook-full.py
=======================
Parses every column of the "Ranked Listings" sheet from
``data/uploads/socal_apartment_ranked_research_workbook.xlsx`` and writes:
- src/lib/data/generated/workbook-listings-full.json
- src/lib/data/generated/workbook-fields-manifest.json

Use this when you change the .xlsx and want the inline workbook table at
/apartments/workbook to reflect it. Run via:

    npm run import:workbook-full

This complements ``import-ranked-workbook.py`` which produces the
link-resolved + geocoded JSON that powers the apply engine.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - friendly fallback
    print("openpyxl is required. Install it via:  pip install openpyxl", file=sys.stderr)
    sys.exit(2)

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "uploads" / "socal_apartment_ranked_research_workbook.xlsx"
OUT_LISTINGS = ROOT / "src" / "lib" / "data" / "generated" / "workbook-listings-full.json"
OUT_MANIFEST = ROOT / "src" / "lib" / "data" / "generated" / "workbook-fields-manifest.json"
OUT_LISTINGS.parent.mkdir(parents=True, exist_ok=True)


def slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(s).lower()).strip("_")


def slug_id(name: str, city: str, rank: int | str) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", f"{name} {city}".lower()).strip("-")
    return f"wb-{rank}-{base}"


NUMERIC_KEYS = {
    "rank",
    "match_percent",
    "rent_low",
    "rent_high",
    "budget_fit_score",
    "safety_score",
    "traffic_freeway_score",
    "screening_difficulty",
    "screening_ease_pts",
    "app_move_in_pts",
    "match",
}


def main() -> int:
    if not XLSX.exists():
        print(f"workbook not found: {XLSX}", file=sys.stderr)
        return 1
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    if "Ranked Listings" not in wb.sheetnames:
        print(f"sheet 'Ranked Listings' missing — found {wb.sheetnames}", file=sys.stderr)
        return 1

    ws = wb["Ranked Listings"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("workbook is empty", file=sys.stderr)
        return 1

    headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(rows[0])]
    keys = [slug(h) for h in headers]

    listings: list[dict] = []
    for r in rows[1:]:
        if not r or r[0] is None:
            continue
        rec: dict = {keys[i]: r[i] for i in range(min(len(keys), len(r)))}
        for k in NUMERIC_KEYS & set(rec.keys()):
            v = rec.get(k)
            if isinstance(v, str):
                try:
                    rec[k] = float(v)
                except ValueError:
                    pass
        rec["id"] = slug_id(
            rec.get("property_name") or "",
            rec.get("city") or "",
            rec.get("rank") or len(listings) + 1,
        )
        listings.append(rec)

    OUT_LISTINGS.write_text(
        json.dumps(listings, indent=2, ensure_ascii=False, default=str),
        encoding="utf-8",
    )
    OUT_MANIFEST.write_text(
        json.dumps(
            [{"key": k, "label": h} for k, h in zip(keys, headers)],
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    print(f"wrote {len(listings)} listings → {OUT_LISTINGS}")
    print(f"wrote field manifest ({len(keys)} columns) → {OUT_MANIFEST}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
